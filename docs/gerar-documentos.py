#!/usr/bin/env python3
"""Converte os documentos Markdown do projeto em .docx formatado.

Fonte única: o conteúdo vive nos arquivos .md em docs/. Este script apenas
deriva o .docx — nunca o contrário. Editar o .docx e não o .md é como a
divergência começa.

    .venv-slides/bin/python docs/gerar-documentos.py

Saída, na raiz do projeto:
    Documento_de_Visao_Achei_um_Bicho.docx
    Cronograma_Achei_um_Bicho.docx
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
DOCS = RAIZ / "docs"

TAPAJOS = RGBColor(0x12, 0x3E, 0x4A)
PERDI = RGBColor(0x2F, 0x85, 0x87)
APAGADO = RGBColor(0x5F, 0x70, 0x76)
TEXTO = RGBColor(0x1A, 0x1A, 0x1A)

FONTE = "Calibri"
FONTE_MONO = "Consolas"

DOCUMENTOS = [
    {
        "md": "documento-de-visao.md",
        "docx": "Documento_de_Visao_Achei_um_Bicho.docx",
        "titulo": "Documento de Visão",
        "subtitulo": "Achei um Bicho — mural de cães e gatos perdidos e encontrados em Santarém-PA",
    },
    {
        "md": "cronograma.md",
        "docx": "Cronograma_Achei_um_Bicho.docx",
        "titulo": "Cronograma do Projeto",
        "subtitulo": "Achei um Bicho — planejamento de 09/09/2026 a 21/11/2026",
    },
]

EMOJIS = {"🔴": "", "🔒": "[restrito]", "✓": "-", "•": "-"}


# ------------------------------------------------------------------ utilidades

def limpar(texto):
    for emoji, troca in EMOJIS.items():
        texto = texto.replace(emoji, troca)
    return texto.strip()


def sombrear(celula, cor_hex):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(sombra)


def campo(paragrafo, instrucao):
    """Insere um campo do Word (PAGE, NUMPAGES, TOC)."""
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instrucao
    separa = OxmlElement("w:fldChar")
    separa.set(qn("w:fldCharType"), "separate")
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    run = paragrafo.add_run()
    for elemento in (inicio, texto, separa, fim):
        run._r.append(elemento)


def escrever_inline(paragrafo, texto, tamanho=10.5, cor=TEXTO):
    """Interpreta **negrito**, `código` e [texto](url) dentro de uma linha."""
    texto = limpar(texto)
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = paragrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = paragrafo.add_run(parte[1:-1])
            run.italic = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = paragrafo.add_run(parte[1:-1])
            run.font.name = FONTE_MONO
            run.font.size = Pt(tamanho - 1)
        else:
            run = paragrafo.add_run(parte)
            run.font.name = FONTE
        run.font.size = Pt(tamanho)
        run.font.color.rgb = cor
    return paragrafo


# ------------------------------------------------------------------ estrutura

def configurar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamanho in ((1, 17), (2, 13.5), (3, 11.5)):
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = FONTE
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = TAPAJOS
        estilo.paragraph_format.space_before = Pt(16 if nivel == 1 else 12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True


def montar_capa(doc, titulo, subtitulo, metadados, secoes):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("UNIVERSIDADE FEDERAL DO OESTE DO PARÁ")
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = APAGADO

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Laboratório de Desenvolvimento de Software")
    run.font.size = Pt(11)
    run.font.color.rgb = APAGADO

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(titulo)
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = TAPAJOS

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitulo)
    run.font.size = Pt(13)
    run.font.color.rgb = PERDI

    for _ in range(5):
        doc.add_paragraph()

    if metadados:
        tabela = doc.add_table(rows=0, cols=2)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        for chave, valor in metadados:
            linha = tabela.add_row()
            p_chave = linha.cells[0].paragraphs[0]
            p_chave.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p_chave.add_run(f"{chave}:")
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = APAGADO
            escrever_inline(linha.cells[1].paragraphs[0], valor)
        for linha in tabela.rows:
            linha.cells[0].width = Cm(4.5)
            linha.cells[1].width = Cm(10)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("Sumário", level=1)
    for nivel, texto_secao in secoes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0 if nivel == 2 else 0.8)
        run = p.add_run(limpar(texto_secao))
        run.font.size = Pt(11 if nivel == 2 else 10)
        run.font.bold = nivel == 2
        run.font.color.rgb = TAPAJOS if nivel == 2 else APAGADO
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def montar_rodape(doc, titulo):
    secao = doc.sections[0]
    secao.top_margin = secao.bottom_margin = Cm(2.5)
    secao.left_margin = secao.right_margin = Cm(2.5)

    cabecalho = secao.header.paragraphs[0]
    cabecalho.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = cabecalho.add_run(f"Achei um Bicho · {titulo}")
    run.font.size = Pt(8.5)
    run.font.color.rgb = APAGADO

    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rodape.add_run("Página ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = APAGADO
    campo(rodape, "PAGE")
    run = rodape.add_run(" de ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = APAGADO
    campo(rodape, "NUMPAGES")


def montar_tabela(doc, linhas_md):
    celulas = [
        [c.strip() for c in linha.strip().strip("|").split("|")]
        for linha in linhas_md
        if not re.match(r"^\|[\s:|-]+\|$", linha.strip())
    ]
    if not celulas:
        return
    colunas = max(len(linha) for linha in celulas)
    tabela = doc.add_table(rows=0, cols=colunas)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER

    # largura proporcional ao conteúdo: sem isso o Word divide igual e a coluna
    # de texto longo fica espremida ao lado de uma coluna com dois caracteres
    util = 16.0  # cm entre as margens
    pesos = []
    for j in range(colunas):
        textos = [linha[j] if j < len(linha) else "" for linha in celulas]
        # o teto impede que uma coluna de texto longo esmague as curtas; o piso
        # pela maior palavra impede que a palavra quebre no meio
        media = min(sum(len(t) for t in textos) / len(textos), 48)
        maior_palavra = max((len(w) for t in textos for w in t.split()), default=4)
        pesos.append(max(4, media, maior_palavra * 1.4))
    total = sum(pesos)
    larguras = [max(1.4, util * peso / total) for peso in pesos]
    excedente = sum(larguras) - util
    if excedente > 0:
        maior = larguras.index(max(larguras))
        larguras[maior] -= excedente
    tabela.autofit = False
    # sem tblLayout fixed, Word e Pages ignoram a largura das células
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tabela._tbl.tblPr.append(layout)

    for coluna, largura in zip(tabela.columns, larguras):
        coluna.width = Cm(largura)

    for i, conteudo in enumerate(celulas):
        linha = tabela.add_row()
        for j in range(colunas):
            celula = linha.cells[j]
            celula.paragraphs[0].paragraph_format.space_after = Pt(2)
            valor = conteudo[j] if j < len(conteudo) else ""
            celula.width = Cm(larguras[j])
            escrever_inline(celula.paragraphs[0], valor, tamanho=9.5)
            if i == 0:
                sombrear(celula, "123E4A")
                for run in celula.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif i % 2 == 0:
                sombrear(celula, "F2F5F2")
    doc.add_paragraph()


def converter(md, doc):
    linhas = md.split("\n")
    i = 0
    buffer_tabela = []
    dentro_codigo = False
    buffer_codigo = []

    def descarregar_tabela():
        nonlocal buffer_tabela
        if buffer_tabela:
            montar_tabela(doc, buffer_tabela)
            buffer_tabela = []

    while i < len(linhas):
        linha = linhas[i]
        despida = linha.strip()

        if despida.startswith("```"):
            if dentro_codigo:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                run = p.add_run("\n".join(buffer_codigo))
                run.font.name = FONTE_MONO
                run.font.size = Pt(8.5)
                run.font.color.rgb = TAPAJOS
                buffer_codigo = []
            dentro_codigo = not dentro_codigo
            i += 1
            continue
        if dentro_codigo:
            buffer_codigo.append(linha)
            i += 1
            continue

        figura = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", despida)
        if figura:
            caminho = DOCS / figura.group(2)
            if caminho.exists():
                doc.add_picture(str(caminho), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                legenda = figura.group(1)
                if legenda:
                    p_leg = doc.add_paragraph()
                    p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p_leg.add_run(legenda)
                    run.font.size = Pt(9)
                    run.font.italic = True
                    run.font.color.rgb = APAGADO
            i += 1
            continue

        if despida.startswith("|"):
            buffer_tabela.append(linha)
            i += 1
            continue
        descarregar_tabela()

        if not despida or despida == "---":
            i += 1
            continue

        if despida.startswith("#"):
            nivel = len(despida) - len(despida.lstrip("#"))
            titulo = limpar(despida.lstrip("#").strip())
            if nivel == 1:
                i += 1
                continue  # o título vira capa
            doc.add_heading(titulo, level=min(nivel - 1, 3))
            i += 1
            continue

        if despida.startswith(">"):
            conteudo = limpar(despida.lstrip(">").strip())
            if not conteudo:
                i += 1
                continue
            if conteudo.startswith("#"):
                conteudo = conteudo.lstrip("#").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            escrever_inline(p, conteudo, cor=TAPAJOS)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        item = re.match(r"^([-*]|\d+\.)\s+(.*)$", despida)
        if item:
            # linha indentada logo abaixo é continuação do mesmo item, não
            # parágrafo novo
            pedacos = [item.group(2)]
            j = i + 1
            while j < len(linhas):
                seguinte = linhas[j]
                if not seguinte.strip() or not seguinte.startswith((" ", "\t")):
                    break
                if re.match(r"^([-*]|\d+\.)\s+", seguinte.strip()):
                    break
                pedacos.append(seguinte.strip())
                j += 1
            estilo = "List Number" if item.group(1).endswith(".") else "List Bullet"
            escrever_inline(doc.add_paragraph(style=estilo), " ".join(pedacos))
            i = j
            continue

        # parágrafo: junta linhas seguidas
        paragrafo = [despida]
        j = i + 1
        while j < len(linhas):
            seguinte = linhas[j].strip()
            if (not seguinte or seguinte.startswith(("#", "|", ">", "---", "```"))
                    or re.match(r"^([-*]|\d+\.)\s+", seguinte)):
                break
            paragrafo.append(seguinte)
            j += 1
        escrever_inline(doc.add_paragraph(), " ".join(paragrafo))
        i = j

    descarregar_tabela()


def extrair_secoes(md):
    """Títulos de nível 2 e 3 do Markdown, na ordem, para o sumário."""
    secoes = []
    dentro_codigo = False
    for linha in md.split("\n"):
        if linha.strip().startswith("```"):
            dentro_codigo = not dentro_codigo
        if dentro_codigo:
            continue
        achado = re.match(r"^(#{2,3})\s+(.+)$", linha.strip())
        if achado:
            secoes.append((len(achado.group(1)), achado.group(2).strip()))
    return secoes


def extrair_metadados(md):
    """Lê as linhas '**Chave:** valor' antes do primeiro separador."""
    metadados = []
    for linha in md.split("\n"):
        if linha.strip() == "---":
            break
        achado = re.match(r"^\*\*([^:*]+):\*\*\s*(.+)$", linha.strip())
        if achado:
            metadados.append((achado.group(1).strip(), achado.group(2).strip()))
    return metadados


def corpo_sem_metadados(md):
    linhas = md.split("\n")
    for i, linha in enumerate(linhas):
        if linha.strip() == "---":
            return "\n".join(linhas[i + 1:])
    return md


def main():
    for spec in DOCUMENTOS:
        origem = DOCS / spec["md"]
        md = origem.read_text(encoding="utf-8")

        doc = Document()
        configurar_estilos(doc)
        montar_rodape(doc, spec["titulo"])
        corpo = corpo_sem_metadados(md)
        montar_capa(doc, spec["titulo"], spec["subtitulo"],
                    extrair_metadados(md), extrair_secoes(corpo))
        converter(corpo, doc)

        destino = RAIZ / spec["docx"]
        doc.save(str(destino))
        print(f"{destino.name}  ←  docs/{spec['md']}")


if __name__ == "__main__":
    main()
