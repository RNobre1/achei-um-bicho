#!/usr/bin/env python3
"""Gera o deck da apresentação do alpha do Achei um Bicho (16/10/2026).

Reaproveita o estilo do deck do kickoff (docs/gerar-slides.py): mesmas cores,
fontes e peças. As telas, em docs/alpha/telas/, são prints do APK num celular Android
e da versão web (o mesmo aplicativo), com os casos de demonstração no ar. Caso
aberto, aviso, pistas e meus casos vêm da web: precisam do GPS no campus ou da
versão 0.1.4, e o simulador de posição do Android não alcança a localização do
Google Play Services.

    .venv-slides/bin/python docs/gerar-slides-alpha.py

Saída: Achei_um_Bicho_Alpha.pptx na raiz do projeto.
"""

import importlib.util
from pathlib import Path

import segno
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

RAIZ = Path(__file__).resolve().parent.parent
ALPHA = RAIZ / "docs" / "alpha"
TELAS = ALPHA / "telas"
UML = RAIZ / "docs" / "uml"
SAIDA = RAIZ / "Achei_um_Bicho_Alpha.pptx"

_spec = importlib.util.spec_from_file_location("kickoff", RAIZ / "docs" / "gerar-slides.py")
k = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(k)

LINK_WEB = "https://achei-um-bicho.onrender.com"
# release no GitHub: link permanente e público, ao contrário do artefato do Expo, que
# expira. "latest" aponta sempre para a release mais nova: versão nova não muda o QR
LINK_APK = "https://github.com/RNobre1/achei-um-bicho/releases/latest/download/achei-um-bicho-alpha.apk"

# Contagem da suíte completa no fechamento do alpha, por camada, tirada da integração
# contínua de 15/10 (execução 38798844785, commit 5947caa). Conferir de novo a cada deck.
TESTES = [
    ("Contratos e vocabulário", 47),
    ("API, unidade", 75),
    ("API, integração com Postgres real", 95),
    ("Aplicativo, regras de tela", 99),
]

STATUS = [
    ("F1", "Publicar bicho perdido", True),
    ("F2", "Publicar bicho encontrado, sem conta", True),
    ("F3", "Mural por proximidade", True),
    ("F4", "Caso aberto e avistamento", True),
    ("F5", "Pareamento determinístico", True),
    ("F6", "Verificação por sinal particular", False),
    ("F7", "Contato mediado", False),
    ("F8", "Ciclo de vida do caso", False),
    ("F9", "Publicação sem conexão", False),
]

TELAS_DECK = [
    dict(arquivo="01-casos-perto.png", titulo="Casos perto de você",
         resolve="O mural abre em volta de quem está olhando, com raio de 1,6 km.",
         pontos=[
             "A consulta por raio roda no banco, com índice espacial: ST_DWithin no PostGIS.",
             "Mural vazio não mostra caso fictício: diz quantos existem até 5 km e oferece ampliar ou convidar vizinhos.",
         ], fs=["F3"]),
    dict(arquivo="02-publicar.png", extra="05-codigo.png", titulo="Publicar em quatro passos",
         resolve="Quem perdeu e quem achou publicam pelo mesmo caminho. Quem achou não precisa de conta.",
         pontos=[
             "Porte, cor e coleira são escolha fechada: exatamente o que o pareamento compara.",
             "No fim sai um código de edição, guardado no celular; o banco só guarda o hash dele.",
         ], fs=["F1", "F2"]),
    dict(arquivo="03-caso-aberto.png", extra="04-avistamento.png", titulo="Caso aberto e avistamento",
         resolve="Quem passou na rua avisa sem conta; quem publicou vê onde e quando.",
         pontos=[
             "O mapa mostra uma área aproximada, sempre o mesmo círculo; tocar abre no Google Maps, sem pino.",
             "Quem publicou vê a distância e o bairro de quem avisou, nunca o ponto exato.",
         ], fs=["F4"]),
    dict(arquivo="06-pistas.png", extra="07-meus-casos.png", titulo="Pode ser o mesmo bicho",
         resolve="Ao publicar, o sistema procura do outro lado e mostra o que bateu e o que não bateu.",
         pontos=[
             "Sem nota nem porcentagem: o par que bate em tudo vem primeiro, mesmo que outro esteja mais perto.",
             "Quem decide é uma pessoa. “Não é” descarta o par para os dois lados.",
         ], fs=["F5"]),
]


def qr(link, arquivo):
    caminho = ALPHA / arquivo
    segno.make(link, error="m").save(str(caminho), scale=12, border=2, dark="#123E4A", light="#FFFFFF")
    return caminho


def slide_capa(prs):
    s = k.novo_slide(prs, fundo=k.TAPAJOS)
    k.retangulo(s, esq=0, topo=0, larg=4.6, alt=k.A, cor=k.RGBColor(0x0D, 0x31, 0x3B))
    k.texto(s, "LABORATÓRIO DE DESENVOLVIMENTO DE SOFTWARE · UFOPA", esq=1.0, topo=1.5, larg=11,
            alt=0.3, tam=11, fonte=k.DADO_FONTE, cor=k.RGBColor(0x8F, 0xB4, 0xB2), entreletras=1.4)
    k.texto(s, "Achei um Bicho", esq=0.95, topo=1.95, larg=11, alt=1.5, tam=66,
            fonte=k.TITULO_FONTE, cor=k.CHUVA, negrito=True, espaco=0.92)
    k.texto(s, "Alpha: publicar, ver, avisar e aproximar, funcionando", esq=1.0, topo=3.35, larg=10,
            alt=0.6, tam=22, cor=k.RGBColor(0xB6, 0xCE, 0xCB))
    k.costura(s, esq=1.0, topo=4.35, larg=3.4)
    k.texto(s, ["Desafio 2 — Causa Animal e Cultura · ODS 15 (Vida Terrestre)",
                "Alpha release · 16 de outubro de 2026"],
            esq=1.0, topo=4.75, larg=9, alt=1.0, tam=14, cor=k.RGBColor(0x9F, 0xBD, 0xBB), espaco=1.3)
    k.texto(s, "EQUIPE", esq=1.0, topo=5.85, larg=9, alt=0.3, tam=10, fonte=k.DADO_FONTE,
            cor=k.RGBColor(0x6F, 0x97, 0x99), entreletras=1.4)
    k.texto(s, "Bruna Kellen · Murilo Falcão · Rafael Nobre", esq=1.0, topo=6.15, larg=9, alt=0.4,
            tam=17, cor=k.RGBColor(0xB6, 0xCE, 0xCB))


def slide_onde_estamos(prs, numero, qr_web):
    s = k.novo_slide(prs)
    k.texto(s, "ONDE ESTAMOS", esq=0.95, topo=0.72, larg=8, alt=0.3, tam=11, fonte=k.DADO_FONTE,
            cor=k.PERDI, entreletras=1.4)
    k.texto(s, "Cinco de nove, no ar", esq=0.9, topo=1.12, larg=8.5, alt=0.8, tam=38,
            fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True)
    for i, (codigo, nome, pronto) in enumerate(STATUS):
        y = 2.2 + i * 0.47
        k.etiqueta(s, codigo, esq=0.9, topo=y, cor_fundo=k.PERDI if pronto else k.ACHEI,
                   cor_texto=k.BRANCO if pronto else k.TAPAJOS, larg=0.5)
        k.texto(s, nome, esq=1.6, topo=y + 0.02, larg=5.0, alt=0.3, tam=15,
                cor=k.TAPAJOS if pronto else k.APAGADO, negrito=pronto)
        k.texto(s, "pronto" if pronto else "próximas sprints", esq=6.6, topo=y + 0.05, larg=2.2,
                alt=0.3, tam=11, fonte=k.DADO_FONTE, cor=k.PERDI if pronto else k.APAGADO)

    k.retangulo(s, esq=9.5, topo=2.05, larg=2.95, alt=4.35, cor=k.TAPAJOS)
    k.texto(s, "ABRA AGORA", esq=9.8, topo=2.35, larg=2.4, alt=0.3, tam=10, fonte=k.DADO_FONTE,
            cor=k.RGBColor(0x8F, 0xB4, 0xB2), entreletras=1.4)
    k.texto(s, "A versão web está no ar", esq=9.8, topo=2.68, larg=2.45, alt=0.8, tam=19,
            fonte=k.TITULO_FONTE, cor=k.CHUVA, negrito=True, espaco=1.05)
    s.shapes.add_picture(str(qr_web), Inches(10.12), Inches(3.75), height=Inches(1.75))
    k.texto(s, "achei-um-bicho.onrender.com", esq=9.75, topo=5.65, larg=2.6, alt=0.4, tam=11,
            fonte=k.DADO_FONTE, cor=k.RGBColor(0xB6, 0xCE, 0xCB), alinha=PP_ALIGN.CENTER)
    k.rodape(s, numero)


def slide_tela(prs, dados, numero):
    s = k.novo_slide(prs)
    altura = 6.26
    largura = altura * 390 / 844
    s.shapes.add_picture(str(TELAS / dados["arquivo"]), Inches(0.75), Inches(0.62), height=Inches(altura))
    x = 0.75 + largura + 0.45
    if dados.get("extra"):
        s.shapes.add_picture(str(TELAS / dados["extra"]), Inches(x), Inches(0.62), height=Inches(altura))
        x += largura + 0.55
    else:
        x += 0.25
    larg_texto = k.L - x - 0.7
    k.texto(s, dados["titulo"], esq=x - 0.05, topo=1.1, larg=larg_texto, alt=1.4, tam=34,
            fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True, espaco=1.02)
    k.texto(s, dados["resolve"], esq=x, topo=2.55, larg=larg_texto, alt=1.2, tam=17, cor=k.TEXTO, espaco=1.35)
    y = 3.95
    for ponto in dados["pontos"]:
        k.retangulo(s, esq=x, topo=y + 0.12, larg=0.16, alt=0.045, cor=k.ACHEI)
        k.texto(s, ponto, esq=x + 0.38, topo=y, larg=larg_texto - 0.4, alt=0.9, tam=13.5,
                cor=k.APAGADO, espaco=1.35)
        y += 0.95
    k.texto(s, "ATENDE", esq=x, topo=5.95, larg=2, alt=0.3, tam=10, fonte=k.DADO_FONTE,
            cor=k.APAGADO, entreletras=1.4)
    for i, codigo in enumerate(dados["fs"]):
        k.etiqueta(s, codigo, esq=x + i * 0.72, topo=6.28, cor_fundo=k.TAPAJOS, larg=0.62)
    k.rodape(s, numero)


def caixa_arquitetura(s, rotulo, detalhe, esq, cor):
    k.retangulo(s, esq=esq, topo=2.2, larg=3.3, alt=1.25, cor=cor)
    k.texto(s, rotulo, esq=esq + 0.25, topo=2.38, larg=2.9, alt=0.4, tam=17, fonte=k.TITULO_FONTE,
            cor=k.CHUVA if cor == k.TAPAJOS else k.TAPAJOS, negrito=True)
    k.texto(s, detalhe, esq=esq + 0.25, topo=2.82, larg=2.9, alt=0.5, tam=11.5, fonte=k.DADO_FONTE,
            cor=k.RGBColor(0xB6, 0xCE, 0xCB) if cor == k.TAPAJOS else k.TAPAJOS)


def slide_por_dentro(prs, numero):
    s = k.novo_slide(prs)
    k.texto(s, "POR DENTRO", esq=0.95, topo=0.72, larg=8, alt=0.3, tam=11, fonte=k.DADO_FONTE,
            cor=k.PERDI, entreletras=1.4)
    k.texto(s, "Três peças, todas grátis", esq=0.9, topo=1.12, larg=11, alt=0.8, tam=38,
            fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True)
    caixa_arquitetura(s, "Aplicativo", "React Native + Expo\nAPK pelo EAS · web no Render", 0.9, k.RGBColor(0xE4, 0xE9, 0xE4))
    caixa_arquitetura(s, "API", "NestJS + Prisma\nDocker no Render", 5.0, k.TAPAJOS)
    caixa_arquitetura(s, "Banco", "PostgreSQL 16 + PostGIS\nNeon", 9.1, k.RGBColor(0xE4, 0xE9, 0xE4))
    for x in (4.3, 8.4):
        k.costura(s, esq=x, topo=2.82, larg=0.6)
    pontos = [
        ("Raio no banco", "ST_DWithin usa o índice GIST; trocar pela distância no WHERE vira varredura da tabela."),
        ("Pareamento explicado", "Compara porte, cor, coleira, distância e tempo. Mostra o que bateu e o que não."),
        ("Privacidade por construção", "Área aproximada fixa por caso, foto sem EXIF, código de edição guardado como hash."),
        ("Uma fonte só", "Vocabulário e formato das requisições declarados uma vez; banco, API e telas derivam dali."),
    ]
    for i, (titulo, detalhe) in enumerate(pontos):
        coluna, linha = divmod(i, 2)
        x = 0.9 + coluna * 6.0
        y = 4.05 + linha * 1.25
        k.retangulo(s, esq=x, topo=y + 0.12, larg=0.16, alt=0.045, cor=k.ALTER)
        k.texto(s, titulo, esq=x + 0.38, topo=y, larg=5.3, alt=0.35, tam=15, cor=k.TAPAJOS, negrito=True)
        k.texto(s, detalhe, esq=x + 0.38, topo=y + 0.36, larg=5.3, alt=0.7, tam=12.5, cor=k.APAGADO, espaco=1.3)
    k.rodape(s, numero)


def slide_modelagem(prs, numero):
    s = k.novo_slide(prs)
    k.texto(s, "MODELAGEM", esq=0.95, topo=0.72, larg=8, alt=0.3, tam=11, fonte=k.DADO_FONTE,
            cor=k.PERDI, entreletras=1.4)
    k.texto(s, "UML gerada do próprio código", esq=0.9, topo=1.12, larg=11.5, alt=0.8, tam=38,
            fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True)
    diagramas = [
        ("casos-de-uso.png", "Casos de uso"),
        ("classes.png", "Classes, do schema do banco"),
        ("estados.png", "Estados do caso, das transições declaradas"),
        ("sequencia-publicar.png", "Sequência: publicar e receber pistas"),
    ]
    for i, (arquivo, legenda) in enumerate(diagramas):
        coluna, linha = i % 2, i // 2
        x = 0.9 + coluna * 6.0
        y = 2.05 + linha * 2.45
        caminho = UML / arquivo
        if caminho.exists():
            s.shapes.add_picture(str(caminho), Inches(x), Inches(y), height=Inches(1.85))
        k.texto(s, legenda, esq=x, topo=y + 1.92, larg=5.6, alt=0.3, tam=11.5, fonte=k.DADO_FONTE, cor=k.APAGADO)
    k.rodape(s, numero)


def slide_qualidade(prs, numero):
    s = k.novo_slide(prs)
    total = sum(n for _, n in TESTES)
    k.texto(s, "QUALIDADE", esq=0.95, topo=0.72, larg=8, alt=0.3, tam=11, fonte=k.DADO_FONTE,
            cor=k.PERDI, entreletras=1.4)
    k.texto(s, f"{total} testes automáticos", esq=0.9, topo=1.12, larg=11, alt=0.8, tam=38,
            fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True)
    for i, (camada, n) in enumerate(TESTES):
        x = 0.9 + i * 2.95
        k.retangulo(s, esq=x, topo=2.35, larg=2.7, alt=0.045, cor=[k.PERDI, k.TAPAJOS, k.ALTER, k.ACHEI][i])
        k.texto(s, str(n), esq=x, topo=2.55, larg=2.7, alt=1.0, tam=48, fonte=k.TITULO_FONTE, cor=k.TAPAJOS, negrito=True)
        k.texto(s, camada, esq=x, topo=3.55, larg=2.6, alt=0.8, tam=13.5, cor=k.APAGADO, espaco=1.3)
    k.texto(s, [
        "Escrito antes do código, em todo incremento: o teste é visto falhando antes de passar.",
        "A integração roda contra PostgreSQL com PostGIS de verdade, num contêiner, e não contra imitação.",
        "Diagramas e restrições do banco gerados do código têm teste que reprova quando ficam para trás.",
        "Testes que conferem, a cada rodada, que contato, código e ponto exato nunca saem nas respostas públicas.",
    ], esq=0.9, topo=4.6, larg=11.5, alt=2.0, tam=14, cor=k.TEXTO, espaco=1.3)
    k.rodape(s, numero)


def slide_fecho(prs, qr_web, qr_apk):
    s = k.novo_slide(prs, fundo=k.TAPAJOS)
    k.retangulo(s, esq=0, topo=0, larg=4.6, alt=k.A, cor=k.RGBColor(0x0D, 0x31, 0x3B))
    k.texto(s, "PRÓXIMOS PASSOS", esq=1.0, topo=0.95, larg=8, alt=0.3, tam=11, fonte=k.DADO_FONTE,
            cor=k.RGBColor(0x8F, 0xB4, 0xB2), entreletras=1.4)
    k.texto(s, [
        "F6 · verificação por sinal particular antes de qualquer contato",
        "F7 · contato liberado só pelo autor, com teste de vazamento",
        "F8 · desfecho obrigatório; caso parado é cobrado e expira",
        "F9 · publicar sem conexão",
        "Primeira rodada com protetores de Santarém",
    ], esq=1.0, topo=1.35, larg=7.2, alt=2.6, tam=15.5, cor=k.RGBColor(0xB6, 0xCE, 0xCB), espaco=1.3)
    k.texto(s, "Obrigado", esq=0.95, topo=4.3, larg=8, alt=1.0, tam=50, fonte=k.TITULO_FONTE,
            cor=k.CHUVA, negrito=True)
    k.costura(s, esq=1.0, topo=5.45, larg=3.4)
    k.texto(s, "Bruna Kellen · Murilo Falcão · Rafael Nobre", esq=1.0, topo=5.85, larg=8, alt=0.4,
            tam=15, cor=k.RGBColor(0x8F, 0xB4, 0xB2))
    for x, imagem, rotulo, nota in ((8.6, qr_web, "Versão web", "qualquer celular"),
                                    (10.75, qr_apk, "APK Android", "instalar sem loja")):
        s.shapes.add_picture(str(imagem), Inches(x + 0.15), Inches(2.6), height=Inches(1.6))
        k.texto(s, rotulo, esq=x, topo=4.35, larg=1.9, alt=0.35, tam=14, cor=k.CHUVA, negrito=True,
                alinha=PP_ALIGN.CENTER)
        k.texto(s, nota, esq=x, topo=4.68, larg=1.9, alt=0.35, tam=10.5, fonte=k.DADO_FONTE,
                cor=k.RGBColor(0x6F, 0x97, 0x99), alinha=PP_ALIGN.CENTER)


def main():
    qr_web = qr(LINK_WEB, "qr-web.png")
    qr_apk = qr(LINK_APK, "qr-apk.png")

    prs = Presentation()
    prs.slide_width = Inches(k.L)
    prs.slide_height = Inches(k.A)

    slide_capa(prs)
    slide_onde_estamos(prs, 2, qr_web)
    for i, dados in enumerate(TELAS_DECK):
        slide_tela(prs, dados, numero=3 + i)
    slide_por_dentro(prs, 3 + len(TELAS_DECK))
    slide_modelagem(prs, 4 + len(TELAS_DECK))
    slide_qualidade(prs, 5 + len(TELAS_DECK))
    slide_fecho(prs, qr_web, qr_apk)

    prs.save(str(SAIDA))
    print(f"{SAIDA.name}: {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    main()
