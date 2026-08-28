#!/usr/bin/env python3
"""Gera o deck de apresentação do MVP do Achei um Bicho.

Fonte única do slide: o conteúdo vive aqui, não no .pptx.
Para mudar um texto, mude neste arquivo e rode de novo:

    .venv-slides/bin/python docs/gerar-slides.py

Saída: Achei_um_Bicho_MVP_e_Prototipo.pptx na raiz do projeto.
"""

from pathlib import Path

import segno
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
TELAS = RAIZ / "prototipo" / "telas"
SAIDA = RAIZ / "Achei_um_Bicho_MVP_e_Prototipo.pptx"
QR_PNG = RAIZ / "docs" / "qr-prototipo.png"

LINK_PROTOTIPO = (
    "https://www.figma.com/proto/ozucUL2Tbschqu31zLceoc/Achei-um-bicho"
    "?node-id=1-562&p=f&t=XjHyy3LTLjuc8o0s-1&scaling=min-zoom"
    "&content-scaling=fixed&page-id=0%3A1&starting-point-node-id=1%3A113"
)

# Paleta — a mesma do protótipo, nomeada pelo que existe na frente da cidade
TAPAJOS = RGBColor(0x12, 0x3E, 0x4A)
PERDI = RGBColor(0x2F, 0x85, 0x87)
ACHEI = RGBColor(0xB6, 0xAA, 0x95)
ALTER = RGBColor(0xE2, 0xA8, 0x3B)
CHUVA = RGBColor(0xEE, 0xF2, 0xEE)
APAGADO = RGBColor(0x5F, 0x70, 0x76)
TEXTO = RGBColor(0x16, 0x32, 0x3A)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)

TITULO_FONTE = "Commissioner"
CORPO_FONTE = "IBM Plex Sans"
DADO_FONTE = "IBM Plex Mono"

L = 13.333  # largura do slide, em polegadas
A = 7.5     # altura


# ---------------------------------------------------------------- utilitários

def texto(slide, conteudo, *, esq, topo, larg, alt=0.6, tam=18, fonte=CORPO_FONTE,
          cor=TEXTO, negrito=False, espaco=1.0, alinha=PP_ALIGN.LEFT,
          maiuscula=False, entreletras=None):
    caixa = slide.shapes.add_textbox(Inches(esq), Inches(topo), Inches(larg), Inches(alt))
    quadro = caixa.text_frame
    quadro.word_wrap = True
    quadro.margin_left = quadro.margin_right = 0
    quadro.margin_top = quadro.margin_bottom = 0
    linhas = conteudo if isinstance(conteudo, list) else [conteudo]
    for i, linha in enumerate(linhas):
        p = quadro.paragraphs[0] if i == 0 else quadro.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        if i:
            p.space_before = Pt(6)
        run = p.add_run()
        run.text = linha.upper() if maiuscula else linha
        f = run.font
        f.name = fonte
        f.size = Pt(tam)
        f.bold = negrito
        f.color.rgb = cor
        if entreletras is not None:
            # espaçamento entre letras não tem API própria no python-pptx
            f._rPr.set("spc", str(int(entreletras * 100)))
    return caixa


def retangulo(slide, *, esq, topo, larg, alt, cor):
    forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(esq), Inches(topo),
                                   Inches(larg), Inches(alt))
    forma.fill.solid()
    forma.fill.fore_color.rgb = cor
    forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def costura(slide, *, esq, topo, larg):
    """A assinatura do projeto: dois traços que se encontram num losango.

    Só aparece onde dois lados do problema se relacionam.
    """
    meio = esq + larg / 2
    no = 0.14
    retangulo(slide, esq=esq, topo=topo, larg=larg / 2 - no, alt=0.028, cor=PERDI)
    retangulo(slide, esq=meio + no, topo=topo, larg=larg / 2 - no, alt=0.028, cor=ACHEI)
    losango = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(meio - no / 2),
                                     Inches(topo - 0.055), Inches(no * 1.4), Inches(no * 1.4))
    losango.fill.solid()
    losango.fill.fore_color.rgb = ALTER
    losango.line.color.rgb = TAPAJOS
    losango.line.width = Pt(0.75)
    losango.shadow.inherit = False


def etiqueta(slide, rotulo, *, esq, topo, cor_fundo, cor_texto=BRANCO, larg=0.62):
    forma = retangulo(slide, esq=esq, topo=topo, larg=larg, alt=0.3, cor=cor_fundo)
    quadro = forma.text_frame
    quadro.margin_left = quadro.margin_right = 0
    quadro.margin_top = quadro.margin_bottom = 0
    quadro.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = quadro.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = rotulo
    run.font.name = DADO_FONTE
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = cor_texto
    return forma


def rodape(slide, numero):
    texto(slide, "Achei um Bicho · UFOPA · Laboratório de Desenvolvimento de Software",
          esq=0.9, topo=6.92, larg=8.0, alt=0.3, tam=9.5, fonte=DADO_FONTE, cor=APAGADO)
    texto(slide, f"{numero:02d}", esq=11.6, topo=6.92, larg=0.85, alt=0.3, tam=9.5,
          fonte=DADO_FONTE, cor=APAGADO, alinha=PP_ALIGN.RIGHT)


def novo_slide(prs, fundo=CHUVA):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    retangulo(slide, esq=0, topo=0, larg=L, alt=A, cor=fundo)
    return slide


# ---------------------------------------------------------------- conteúdo

FUNCIONALIDADES = [
    ("F1", "Publicar bicho perdido", "Tutor informa atributos, local, hora e três sinais particulares."),
    ("F2", "Publicar bicho encontrado", "Sem criar conta. Devolve um código para editar depois."),
    ("F3", "Mural por proximidade", "Raio padrão de 1,6 km, visível e ajustável."),
    ("F4", "Caso aberto e avistamento", "Quem viu na rua registra sem expor contato de ninguém."),
    ("F5", "Pareamento determinístico", "Compara atributos e distância. Mostra o que bateu e o que não."),
    ("F6", "Verificação por sinal particular", "Prova antes do contato, sem botão “é meu”."),
    ("F7", "Contato mediado", "Telefone só aparece depois que o autor libera."),
    ("F8", "Ciclo de vida do caso", "Desfecho obrigatório. Caso parado expira."),
    ("F9", "Publicação sem conexão", "Fila no aparelho, sobe quando a internet volta."),
]

TELAS_DECK = [
    dict(arquivo="tela-01.png", num="01 / 07", titulo="Casos perto de você",
         resolve="O mural abre no bairro de quem está olhando, com raio de 1,6 km.",
         pontos=[
             "Azul da água marca quem sumiu; silte do rio marca quem foi achado. A cor faz o trabalho que um ícone faria.",
             "O raio fica no topo, visível e editável — não escondido dentro de ajustes.",
         ],
         fs=["F3"]),
    dict(arquivo="tela-02.png", num="02 / 07", titulo="Caso aberto",
         resolve="Quem passou na rua tem uma informação. Quem perdeu tem uma reivindicação. São caminhos diferentes.",
         pontos=[
             "Duas ações separadas de propósito: “vi esse bicho na rua” e “acho que esse bicho é meu”.",
             "Telefone e endereço exato não aparecem aqui. O mapa mostra área aproximada, nunca a casa de alguém.",
         ],
         fs=["F4"]),
    dict(arquivo="tela-03.png", num="03 / 07", titulo="Publicar — quem perdeu",
         resolve="O lado do tutor pede o que só ele tem: nome, quando sumiu, onde o bicho costuma andar.",
         pontos=[
             "Aqui nasce a prova usada na verificação: o primeiro dos três sinais particulares.",
             "Ninguém mais vê essa resposta. Sem esta etapa, não existiria verificação nenhuma.",
         ],
         fs=["F1", "F6"]),
    dict(arquivo="tela-04.png", num="04 / 07", titulo="Publicar — quem achou",
         resolve="Publicação em quatro passos, sem criar conta, só com campo de escolha fechada.",
         pontos=[
             "Porte, cor dominante e coleira são exatamente os três atributos que o pareamento compara.",
             "Quem acha um bicho na rua está com pressa e com uma mão só. Exigir cadastro perde o caso.",
         ],
         fs=["F2"]),
    dict(arquivo="tela-05.png", num="05 / 07", titulo="Pistas — pode ser o mesmo bicho",
         resolve="O sistema aproxima dois casos e explica por quê. Quem decide é uma pessoa.",
         pontos=[
             "Sem porcentagem de compatibilidade: mostra o que bateu e também o que não bateu.",
             "O segundo par vem com o botão fraco de propósito — o aplicativo não empurra a conclusão.",
         ],
         fs=["F5"]),
    dict(arquivo="tela-06.png", num="06 / 07", titulo="Confirmar que é seu",
         resolve="Como impedir que alguém leve o animal de outra pessoa.",
         pontos=[
             "Não existe botão “é meu”. Existe prova por sinal particular, guardada por quem achou.",
             "O aplicativo não decide propriedade. Contato só é liberado pelo autor do caso.",
         ],
         fs=["F6", "F7"]),
    dict(arquivo="tela-07.png", num="07 / 07", titulo="Meus casos",
         resolve="Acompanhar os próprios casos e informar o desfecho.",
         pontos=[
             "Caso sem resposta há 9 dias é cobrado; sem interação em 15 dias, sai do mural.",
             "É o detalhe que separa este mural de um grupo de WhatsApp cheio de anúncio morto.",
         ],
         fs=["F8"]),
]

# Fora do deck por decisão de 28/08/2026: os dois estados (bairro sem casos e
# publicação offline) e o slide "o que ficou de fora" continuam no protótipo e em
# docs/mvp-funcionalidades.md, mas não entram nesta apresentação.
ESTADOS_FORA_DO_DECK = ["estado-a.png", "estado-b.png"]

# Origem dos números do slide 3. Fica aqui para não virar achismo em slide.
FONTE_DADOS = (
    "Fonte: Kremer, T. (2021). A New Web-Based Tool for RTO-Focused Animal Shelter "
    "Data Analysis. Frontiers in Veterinary Science. Dados do abrigo público de "
    "Dallas, 2019. Consultado em 26/08/2026."
)
RESSALVA_DADOS = (
    "Estudo norte-americano: ponto de partida, não verdade local. A primeira rodada "
    "com protetores de Santarém vai testar se o raio se confirma aqui."
)


# ---------------------------------------------------------------- slides

def slide_capa(prs):
    slide = novo_slide(prs, fundo=TAPAJOS)
    retangulo(slide, esq=0, topo=0, larg=4.6, alt=A, cor=RGBColor(0x0D, 0x31, 0x3B))

    texto(slide, "LABORATÓRIO DE DESENVOLVIMENTO DE SOFTWARE · UFOPA",
          esq=1.0, topo=1.5, larg=11, alt=0.3, tam=11, fonte=DADO_FONTE,
          cor=RGBColor(0x8F, 0xB4, 0xB2), entreletras=1.4)
    texto(slide, "Achei um Bicho", esq=0.95, topo=1.95, larg=11, alt=1.5, tam=66,
          fonte=TITULO_FONTE, cor=CHUVA, negrito=True, espaco=0.92)
    texto(slide, "Mural de cães e gatos perdidos e encontrados em Santarém",
          esq=1.0, topo=3.35, larg=9, alt=0.6, tam=22, cor=RGBColor(0xB6, 0xCE, 0xCB))
    costura(slide, esq=1.0, topo=4.35, larg=3.4)
    texto(slide, ["Desafio 2 — Causa Animal e Cultura · ODS 15 (Vida Terrestre)",
                  "MVP e protótipo de média fidelidade · 28 de agosto de 2026"],
          esq=1.0, topo=4.75, larg=9, alt=1.0, tam=14, cor=RGBColor(0x9F, 0xBD, 0xBB), espaco=1.3)
    texto(slide, "EQUIPE", esq=1.0, topo=5.85, larg=9, alt=0.3, tam=10,
          fonte=DADO_FONTE, cor=RGBColor(0x6F, 0x97, 0x99), entreletras=1.4)
    texto(slide, "Bruna Kellen · Murilo Falcão · Rafael Nobre",
          esq=1.0, topo=6.15, larg=9, alt=0.4, tam=17, cor=RGBColor(0xB6, 0xCE, 0xCB))


def slide_problema(prs):
    slide = novo_slide(prs)
    texto(slide, "O PROBLEMA", esq=0.95, topo=1.0, larg=6, alt=0.3, tam=11,
          fonte=DADO_FONTE, cor=PERDI, entreletras=1.4)
    texto(slide, ["As duas informações existem.", "Elas é que não se encontram."],
          esq=0.9, topo=1.45, larg=6.9, alt=2.4, tam=34, fonte=TITULO_FONTE,
          cor=TAPAJOS, negrito=True, espaco=1.08)
    texto(slide,
          "Em Santarém, quem perde um cão publica num grupo de WhatsApp. Quem encontra "
          "publica em outro. As duas mensagens são do mesmo dia, da mesma cidade, às vezes "
          "do mesmo bairro — e ninguém as compara.",
          esq=0.9, topo=4.05, larg=6.6, alt=1.6, tam=17, cor=TEXTO, espaco=1.35)

    retangulo(slide, esq=8.0, topo=1.45, larg=4.45, alt=3.9, cor=RGBColor(0xE4, 0xE9, 0xE4))
    retangulo(slide, esq=8.0, topo=1.45, larg=0.06, alt=3.9, cor=ALTER)
    texto(slide, "O CONCORRENTE REAL", esq=8.4, topo=1.85, larg=3.8, alt=0.3, tam=10.5,
          fonte=DADO_FONTE, cor=APAGADO, entreletras=1.4)
    texto(slide, "Não é outro aplicativo. É o grupo de WhatsApp.",
          esq=8.4, topo=2.25, larg=3.75, alt=1.4, tam=25, fonte=TITULO_FONTE,
          cor=TAPAJOS, negrito=True, espaco=1.08)
    texto(slide,
          "Funcionalidade que não seja melhor que “postar no grupo” não entra no MVP. "
          "Foi esse o critério de corte de tudo que vem a seguir.",
          esq=8.4, topo=4.0, larg=3.75, alt=1.2, tam=14, cor=TEXTO, espaco=1.35)
    rodape(slide, 2)


def slide_regra(prs):
    slide = novo_slide(prs)
    texto(slide, "A REGRA QUE ORGANIZA O PRODUTO INTEIRO", esq=0.95, topo=1.0, larg=8,
          alt=0.3, tam=11, fonte=DADO_FONTE, cor=PERDI, entreletras=1.4)
    texto(slide, "O bicho está perto de casa.", esq=0.9, topo=1.45, larg=9, alt=1.0,
          tam=42, fonte=TITULO_FONTE, cor=TAPAJOS, negrito=True)

    for i, (numero, legenda, cor) in enumerate([
        ("70%", "dos cães recuperados estavam a menos de 1,6 km de casa", PERDI),
        ("42%", "estavam a menos de um quarteirão", TAPAJOS),
        ("1,6 km", "é o raio padrão do mural, visível e editável", ALTER),
    ]):
        x = 0.9 + i * 4.05
        retangulo(slide, esq=x, topo=2.75, larg=3.75, alt=0.045, cor=cor)
        texto(slide, numero, esq=x, topo=2.95, larg=3.75, alt=1.1, tam=54,
              fonte=TITULO_FONTE, cor=TAPAJOS, negrito=True)
        texto(slide, legenda, esq=x, topo=4.1, larg=3.6, alt=1.2, tam=15,
              cor=APAGADO, espaco=1.35)

    texto(slide,
          "Alerta de vizinhança supera divulgação para a cidade inteira. Toda decisão de "
          "produto deste projeto sai daí — inclusive as que cortamos.",
          esq=0.9, topo=5.15, larg=11.5, alt=0.6, tam=16, cor=TEXTO, espaco=1.35)

    retangulo(slide, esq=0.9, topo=5.85, larg=11.5, alt=0.012, cor=RGBColor(0xC7, 0xD0, 0xCB))
    texto(slide, [FONTE_DADOS, RESSALVA_DADOS], esq=0.9, topo=5.98, larg=11.5, alt=0.85,
          tam=9.5, cor=APAGADO, espaco=1.25)
    rodape(slide, 3)


def slide_mvp(prs):
    slide = novo_slide(prs)
    texto(slide, "O QUE O MVP FAZ", esq=0.95, topo=0.72, larg=8, alt=0.3, tam=11,
          fonte=DADO_FONTE, cor=PERDI, entreletras=1.4)
    texto(slide, "Nove funcionalidades", esq=0.9, topo=1.12, larg=8, alt=0.8, tam=38,
          fonte=TITULO_FONTE, cor=TAPAJOS, negrito=True)

    topo0 = 2.15
    for i, (codigo, nome, detalhe) in enumerate(FUNCIONALIDADES):
        coluna, linha = divmod(i, 5)
        if coluna == 1 and linha > 3:
            coluna, linha = 1, linha
        x = 0.9 + coluna * 4.55
        y = topo0 + linha * 0.95
        etiqueta(slide, codigo, esq=x, topo=y + 0.04, cor_fundo=PERDI, larg=0.5)
        texto(slide, nome, esq=x + 0.68, topo=y, larg=3.4, alt=0.32, tam=15,
              cor=TAPAJOS, negrito=True)
        texto(slide, detalhe, esq=x + 0.68, topo=y + 0.3, larg=3.35, alt=0.6,
              tam=11.5, cor=APAGADO, espaco=1.22)

    # bloco do protótipo, com QR
    retangulo(slide, esq=9.9, topo=2.05, larg=2.55, alt=4.35, cor=TAPAJOS)
    texto(slide, "VEJA AS TELAS", esq=10.15, topo=2.35, larg=2.1, alt=0.3, tam=10,
          fonte=DADO_FONTE, cor=RGBColor(0x8F, 0xB4, 0xB2), entreletras=1.4)
    texto(slide, "Protótipo navegável", esq=10.15, topo=2.68, larg=2.1, alt=0.8,
          tam=19, fonte=TITULO_FONTE, cor=CHUVA, negrito=True, espaco=1.05)
    if QR_PNG.exists():
        slide.shapes.add_picture(str(QR_PNG), Inches(10.42), Inches(3.72), height=Inches(1.75))
    texto(slide, "Aponte a câmera para abrir no celular.",
          esq=10.15, topo=5.62, larg=2.1, alt=0.7, tam=11,
          cor=RGBColor(0xB6, 0xCE, 0xCB), espaco=1.3)
    rodape(slide, 4)


def slide_tela(prs, dados, numero):
    slide = novo_slide(prs)
    caminho = TELAS / dados["arquivo"]
    if caminho.exists():
        slide.shapes.add_picture(str(caminho), Inches(0.95), Inches(0.62), height=Inches(6.26))

    x = 4.35
    texto(slide, dados["num"], esq=x, topo=1.15, larg=4, alt=0.3, tam=11,
          fonte=DADO_FONTE, cor=PERDI, entreletras=1.4)
    texto(slide, dados["titulo"], esq=x - 0.05, topo=1.5, larg=8.2, alt=1.0, tam=40,
          fonte=TITULO_FONTE, cor=TAPAJOS, negrito=True, espaco=1.02)
    texto(slide, dados["resolve"], esq=x, topo=2.62, larg=7.6, alt=1.0, tam=19,
          cor=TEXTO, espaco=1.35)

    y = 3.95
    for ponto in dados["pontos"]:
        retangulo(slide, esq=x, topo=y + 0.12, larg=0.16, alt=0.045, cor=ACHEI)
        texto(slide, ponto, esq=x + 0.38, topo=y, larg=7.2, alt=0.8, tam=14.5,
              cor=APAGADO, espaco=1.35)
        y += 0.92

    texto(slide, "ATENDE", esq=x, topo=5.95, larg=2, alt=0.3, tam=10,
          fonte=DADO_FONTE, cor=APAGADO, entreletras=1.4)
    for i, codigo in enumerate(dados["fs"]):
        etiqueta(slide, codigo, esq=x + i * 0.72, topo=6.28, cor_fundo=TAPAJOS, larg=0.62)
    rodape(slide, numero)


def slide_fecho(prs):
    slide = novo_slide(prs, fundo=TAPAJOS)
    retangulo(slide, esq=0, topo=0, larg=4.6, alt=A, cor=RGBColor(0x0D, 0x31, 0x3B))

    texto(slide, "Obrigado", esq=0.95, topo=2.35, larg=9, alt=1.3, tam=60,
          fonte=TITULO_FONTE, cor=CHUVA, negrito=True)
    texto(slide, "Quem perdeu e quem achou estão na mesma cidade, no mesmo dia. Falta só se encontrarem.",
          esq=1.0, topo=3.62, larg=7.6, alt=0.9, tam=19, cor=RGBColor(0xB6, 0xCE, 0xCB), espaco=1.35)
    costura(slide, esq=1.0, topo=4.72, larg=3.4)
    texto(slide, "Bruna Kellen · Murilo Falcão · Rafael Nobre",
          esq=1.0, topo=5.15, larg=8, alt=0.4, tam=15, cor=RGBColor(0x8F, 0xB4, 0xB2))

    if QR_PNG.exists():
        slide.shapes.add_picture(str(QR_PNG), Inches(10.55), Inches(2.55), height=Inches(1.7))
    texto(slide, "Protótipo navegável", esq=10.0, topo=4.42, larg=2.8, alt=0.35, tam=14,
          cor=CHUVA, negrito=True, alinha=PP_ALIGN.CENTER)
    texto(slide, "Aponte a câmera", esq=10.0, topo=4.75, larg=2.8, alt=0.35, tam=11,
          fonte=DADO_FONTE, cor=RGBColor(0x6F, 0x97, 0x99), alinha=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- montagem

def gerar_qr():
    qr = segno.make(LINK_PROTOTIPO, error="m")
    qr.save(str(QR_PNG), scale=14, border=2, dark="#123E4A", light="#FFFFFF")
    return QR_PNG


def main():
    gerar_qr()

    prs = Presentation()
    prs.slide_width = Inches(L)
    prs.slide_height = Inches(A)

    slide_capa(prs)
    slide_problema(prs)
    slide_regra(prs)
    slide_mvp(prs)
    for i, dados in enumerate(TELAS_DECK):
        slide_tela(prs, dados, numero=5 + i)
    slide_fecho(prs)

    prs.save(str(SAIDA))
    print(f"{SAIDA.name} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
    print(f"QR gerado em {QR_PNG.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
